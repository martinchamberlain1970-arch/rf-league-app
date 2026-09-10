from pathlib import Path
import shutil
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "public/guides/Rack-and-Frame-Captain-and-Vice-Captain-Guide-2026-27.docx"
DELIVERABLE = ROOT / "deliverables/Rack-and-Frame-Captain-and-Vice-Captain-Guide-2026-27.docx"
SHOTS = ROOT / "public/guides/screenshots"

NAVY = "0F2744"
PALE = "EAF4FA"
BORDER = "D9D9D9"
TEAL = RGBColor(0x0F, 0x76, 0x6E)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd")) or OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def borders(table):
    tblPr = table._tbl.tblPr
    el = tblPr.find(qn("w:tblBorders")) or OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single"); node.set(qn("w:sz"), "6"); node.set(qn("w:color"), BORDER)
        el.append(node)
    tblPr.append(el)

def keep(paragraph, next_=False):
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = next_

def set_cell_text(cell, value, bold=False, color="000000"):
    cell.text = ""
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run(value); r.bold = bold; r.font.name = "Arial"; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_bullet(doc, value):
    p = doc.add_paragraph(style="List Bullet"); p.add_run(value); keep(p)
    return p

def add_step(doc, lead, detail):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(lead + "  "); r.bold = True; r.font.color.rgb = TEAL
    p.add_run(detail); keep(p)

def add_figure_page(doc, title, intro, image, caption):
    doc.add_page_break()
    h = doc.add_heading(title, level=1); keep(h, True)
    p = doc.add_paragraph(intro); keep(p, True)
    pic = doc.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.space_before = Pt(3)
    pic.add_run().add_picture(str(SHOTS / image), width=Inches(6.0)); keep(pic)
    cap = doc.add_paragraph(caption); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.style = doc.styles["Caption"]; keep(cap)

doc = Document()
section = doc.sections[0]
section.page_width = Cm(21); section.page_height = Cm(29.7)
section.top_margin = Cm(1.7); section.bottom_margin = Cm(1.7); section.left_margin = Cm(1.8); section.right_margin = Cm(1.8)

styles = doc.styles
styles["Normal"].font.name = "Arial"; styles["Normal"].font.size = Pt(10.5); styles["Normal"].font.color.rgb = RGBColor(0,0,0)
styles["Normal"].paragraph_format.space_after = Pt(6); styles["Normal"].paragraph_format.line_spacing = 1.08
styles["Title"].font.name = "Arial"; styles["Title"].font.size = Pt(30); styles["Title"].font.bold = True; styles["Title"].font.color.rgb = RGBColor(0,0,0)
for name, size in (("Heading 1",20),("Heading 2",14)):
    styles[name].font.name = "Arial"; styles[name].font.size = Pt(size); styles[name].font.bold = True; styles[name].font.color.rgb = RGBColor(0,0,0)
    styles[name].paragraph_format.space_before = Pt(8); styles[name].paragraph_format.space_after = Pt(7)
styles["Caption"].font.name = "Arial"; styles["Caption"].font.size = Pt(9); styles["Caption"].font.italic = True; styles["Caption"].font.color.rgb = RGBColor(0x47,0x55,0x69)

header = section.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
rh = header.add_run("RACK AND FRAME LEAGUE MANAGER"); rh.bold = True; rh.font.name = "Arial"; rh.font.size = Pt(8); rh.font.color.rgb = TEAL
footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
rf = footer.add_run("Captain and Vice Captain Guide  |  2026 and 2027 season"); rf.font.name="Arial"; rf.font.size=Pt(8); rf.font.color.rgb=RGBColor(0x64,0x74,0x8B)

logo = doc.add_paragraph(); logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
logo.add_run().add_picture(str(ROOT / "public/rf-logo.png"), width=Inches(0.72))
title = doc.add_paragraph(style="Title"); title.add_run("Captain and Vice Captain Guide")
sub = doc.add_paragraph("The complete winter league match night process for team officials"); sub.runs[0].font.size=Pt(15); sub.runs[0].font.color.rgb=TEAL
doc.add_paragraph("Gravesend and District Indoor Games League  |  2026 and 2027 season")

doc.add_heading("What this guide covers", level=1)
doc.add_paragraph("Captains and vice-captains use the same Captain Results journey. This guide explains the normal four-player winter lineup, the two-player and three-player exceptions, agreed proxy entry, frame scoring, breaks and final submission.")

doc.add_heading("Match night quick reference", level=1)
table = doc.add_table(rows=1, cols=3); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
widths=[Cm(1.4),Cm(4.2),Cm(11.0)]
for i,(cell,w,label) in enumerate(zip(table.rows[0].cells,widths,["Step","Action","What to do"])):
    cell.width=w; shade(cell,NAVY); set_cell_text(cell,label,True,"FFFFFF")
rows=[("1","Open the fixture","Open Captain Results and confirm the fixture, teams and date."),("2","Choose entry method","Use the normal home-first process, or enable agreed proxy entry before selecting players."),("3","Complete five frames","Enter four singles and one doubles lineup using the appropriate player-availability rule."),("4","Submit lineups","Home submits first; away confirms unless agreed proxy entry is active."),("5","Record the match","Enter every frame score and all breaks of 30 or more."),("6","Review and submit","The home team normally submits the agreed result for league-officer approval.")]
for n,a,b in rows:
    cells=table.add_row().cells
    for i,(cell,w,val) in enumerate(zip(cells,widths,[n,a,b])):
        cell.width=w; shade(cell,"FFFFFF" if int(n)%2 else PALE); set_cell_text(cell,val,bold=i==1)
borders(table)

doc.add_heading("Before selecting players", level=1)
add_bullet(doc,"The digital lineup opens on the day of the fixture. Home normally submits first and away then confirms.")
add_bullet(doc,"If both teams agree that one official will operate the app, enable agreed proxy entry before selecting any players.")
add_bullet(doc,"There is no paper-record option inside this digital journey. If the app cannot be used, keep an accurate manual scorecard and contact the League Secretary for the separate result-upload link.")

add_figure_page(doc,"Normal four player winter lineup","Select four different singles players. Then choose two eligible players for Frame 5 doubles. The action buttons are deliberately placed after the final frame.","captain-winter-four-player.png","Normal winter lineup with four singles, one doubles frame and the lineup actions at the end of the journey.")
add_figure_page(doc,"Agreed proxy entry","Use this only when both teams agree. Enable it at the top of the page before entering lineups; it unlocks both teams' fields so one captain or vice-captain can complete both sides.","captain-proxy-entry.png","Agreed proxy entry opens the opponent fields. The acting official must confirm both lineups and the final result with both teams.")
add_figure_page(doc,"Two player winter lineup","Choose two different players in Frames 1 and 2, then select No Show in Frame 3. A confirmation explains the consequences before anything changes.","captain-winter-two-player.png","After confirmation, Frame 3 is No Show, the system randomly names the nominated player for Frame 4, and both available players fill the doubles automatically.")
add_figure_page(doc,"Three player winter lineup","Choose three different players in Frames 1 to 3, then choose Nominated player in Frame 4 and acknowledge the confirmation.","captain-winter-three-player.png","The system randomly names the nominated player for Frame 4. The captain then chooses any two of the first three players for the doubles.")

doc.add_page_break()
doc.add_heading("Lineup rules and checks", level=1)
doc.add_heading("Four players available", level=2)
add_bullet(doc,"Use four different players across Frames 1 to 4.")
add_bullet(doc,"Choose the doubles pair in Frame 5 and review the complete lineup before submitting.")
doc.add_heading("Only three players available", level=2)
add_bullet(doc,"Use three different players in Frames 1 to 3.")
add_bullet(doc,"Choose Nominated player in Frame 4 and confirm that only three players are available.")
add_bullet(doc,"The system randomly nominates one of the first three players and displays the player's name followed by nominated player in brackets.")
add_bullet(doc,"Choose any two of those three players for the doubles.")
doc.add_heading("Only two players available", level=2)
add_bullet(doc,"Use two different players in Frames 1 and 2.")
add_bullet(doc,"Choose No Show in Frame 3 and acknowledge the confirmation.")
add_bullet(doc,"The system randomly nominates one of the two players for Frame 4 and automatically places both players in the doubles.")
doc.add_heading("Submitting the lineup", level=2)
add_bullet(doc,"Save lineup draft keeps unfinished work on the device.")
add_bullet(doc,"Submit team to opponent only when the home lineup is final. Away then checks and confirms it.")
add_bullet(doc,"The same player cannot be manually selected in more than one ordinary winter singles frame. A system-nominated Frame 4 is the permitted exception.")

add_figure_page(doc,"Review and submit the result","Once both lineups are locked, enter each completed frame, add every break of 30 or more, and save as you progress. After the final frame, review the full scorecard with both teams.","captain-final-scorecard-complete.png","Check player names, all five frame scores and qualifying breaks before selecting Submit match result.")

doc.add_page_break()
doc.add_heading("Complete match night process", level=1)
add_step(doc,"1  Open Captain Results","Select tonight's fixture and confirm the teams, date and lineup status.")
add_step(doc,"2  Decide whether proxy entry is needed","If both sides agree that one official will operate the app, enable agreed proxy entry before choosing players.")
add_step(doc,"3  Enter the lineups","Follow the four-player, three-player or two-player winter route. Home submits first unless proxy entry is active.")
add_step(doc,"4  Open the scorecard","After both lineups are locked, check the players shown for the first frame.")
add_step(doc,"5  Save every completed frame","Enter both scores, record any break of 30 or more and continue through all five frames.")
add_step(doc,"6  Review with both teams","Check every player, score and break before final submission.")
add_step(doc,"7  Submit the result","The home team normally submits by midnight on the following day. The result enters the league-officer approval queue.")

doc.add_heading("If the app causes a problem", level=1)
add_bullet(doc,"Keep the match moving and retain an accurate manual scorecard. Snooker takes priority.")
add_bullet(doc,"Message the League Secretary and allow time for a reply on match evenings.")
add_bullet(doc,"A separate result-upload link can be supplied if the normal app journey cannot be completed.")
add_bullet(doc,"A manual upload requires league-officer authorisation and therefore updates more slowly than a normal app submission.")

doc.add_heading("Correcting mistakes", level=2)
corrections = doc.add_table(rows=1, cols=1)
corrections.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = corrections.cell(0, 0)
shade(cell, PALE)
cell.text = ""
for value in (
    "Correct errors before submission whenever possible.",
    "After submission, contact the League Secretary, Chairman or Treasurer so the result can be returned or corrected during review.",
    "Use the fixture-date request process when a match needs to move. Do not rely only on an informal message.",
):
    p = cell.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(value)
borders(corrections)

OUT.parent.mkdir(parents=True,exist_ok=True); DELIVERABLE.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT); shutil.copy2(OUT,DELIVERABLE)
print(OUT); print(DELIVERABLE)
